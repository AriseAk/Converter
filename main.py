from flask import Flask, request, send_file, jsonify
from flask import render_template
from pymongo import MongoClient
import gridfs
import os
from werkzeug.utils import secure_filename
from pdf2docx import Converter
from PIL import Image
import io
from dotenv import load_dotenv
from bson.objectid import ObjectId
from docx import Document
from fpdf import FPDF
from datetime import datetime, timezone
from apscheduler.schedulers.background import BackgroundScheduler

app = Flask(__name__)

load_dotenv()
# MongoDB connection
client = MongoClient(os.getenv("MONGO_CLIENT"))  # Change to your MongoDB URI
mongo_db = client["file_converter"]
fs = gridfs.GridFS(mongo_db)

# Allowed file extensions
ALLOWED_EXTENSIONS = {"pdf", "jpg", "png", "svg" ,"jpeg" ,"docx"}
UPLOAD_FOLDER = "uploads/"

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

@app.route("/")
def index():
    return render_template("index.html")

# Helper function to check file extension
def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

# Route: Upload and Convert PDF to Word
@app.route("/convert/pdf-to-word", methods=["POST"])
def pdf_to_word():
    if "file" not in request.files:
        return "<p style='color:red;'>No file uploaded!</p>", 400
    
    file = request.files["file"]
    if file.filename == "":
        return "<p style='color:red;'>No selected file!</p>", 400
    
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_path = os.path.join(UPLOAD_FOLDER, filename)
        file.save(file_path)

        # Convert PDF to Word
        docx_path = file_path.replace(".pdf", ".docx")
        cv = Converter(file_path)
        cv.convert(docx_path)
        cv.close()

        # Store in MongoDB (GridFS)
        with open(docx_path, "rb") as converted_file:
            file_id = fs.put(converted_file, filename=os.path.basename(docx_path))

        # Delete local files
        os.remove(file_path)
        os.remove(docx_path)

        # Store metadata in MongoDB
        mongo_db.files.insert_one({
            "filename": filename,
            "original_format": "pdf",
            "converted_format": "docx",
            "status": "converted",
            "file_id": file_id,
            "uploaded_at": datetime.now(timezone.utc)
        })

        # Return HTML response for HTMX
        return f"<p style='color:green;'>File converted! <a href='/download/{file_id}'>Download Word File</a></p>"
    
    return "<p style='color:red;'>Invalid file type!</p>", 400

class PDF(FPDF):
    def header(self):
        self.set_font("Arial", "B", 12)
        self.cell(0, 10, "Converted PDF", ln=True, align="C")

@app.route("/convert/word-to-pdf", methods=["POST"])
def word_to_pdf():
    if "file" not in request.files:
        return jsonify({"error": "No file part"}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "No selected file"}), 400

    if file and file.filename.endswith(".docx"):
        filename = secure_filename(file.filename)
        file_path = os.path.join(UPLOAD_FOLDER, filename)
        file.save(file_path)

        # Convert DOCX to PDF
        pdf_path = file_path.replace(".docx", ".pdf")
        doc = Document(file_path)
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        for para in doc.paragraphs:
            text = para.text
            text = text.encode("latin-1", "ignore").decode("latin-1")  # Fix encoding issue
            pdf.multi_cell(0, 10, txt=text, align="L")

        pdf.output(pdf_path)

        # Store in MongoDB (GridFS)
        with open(pdf_path, "rb") as converted_file:
            file_id = fs.put(converted_file, filename=os.path.basename(pdf_path))

        # Delete local files
        os.remove(file_path)
        os.remove(pdf_path)

        # Store metadata in MongoDB
        mongo_db.files.insert_one({
            "filename": filename,
            "original_format": "docx",
            "converted_format": "pdf",
            "status": "converted",
            "file_id": file_id,
            "uploaded_at": datetime.now(timezone.utc)
        })

        return f"<p style='color:green;'>File converted! <a href='/download/{file_id}'>Download PDF File</a></p>"

    return "<p style='color:red;'>Invalid file type!</p>", 400

@app.route("/download/<file_id>", methods=["GET"])
def download_file(file_id):
    try:
        file = fs.get(ObjectId(file_id))  # Convert file_id to ObjectId
        return send_file(io.BytesIO(file.read()), 
                         download_name=file.filename, 
                         as_attachment=True)
    except Exception as e:
        return jsonify({"error": "File not found or invalid ID"}), 404

# Route: Display Converted Files
@app.route("/converted-files", methods=["GET"])
def converted_files():
    files = list(mongo_db.files.find({}, {"filename": 1, "file_id": 1}))
    return render_template("converted.html", files=files)

# ====== FIXED JPG TO SVG CONVERSION ======
@app.route("/convert/jpg-svg", methods=["POST"])
def jpg_to_svg():
    if "file" not in request.files:
        return jsonify({"error": "No file part"}), 400
    
    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "No selected file"}), 400
    
    if file and allowed_file(file.filename):
        # ✅ Read file directly from memory
        img = Image.open(io.BytesIO(file.read())).convert("RGBA")
        
        # ✅ Convert to SVG (in memory)
        svg_data = convert_jpg_to_svg(img)
        
        # ✅ Store in MongoDB (directly from memory)
        file_id = fs.put(svg_data.encode("utf-8"), filename=file.filename.replace(".jpg", ".svg"))
        
        # ✅ Save metadata
        mongo_db.files.insert_one({
            "filename": file.filename,
            "original_format": "jpg",
            "converted_format": "svg",
            "status": "converted",
            "file_id": file_id
        })

        return f"<p style='color:green;'>File converted! <a href='/download/{file_id}'>Download SVG File</a></p>"

    return "<p style='color:red;'>Invalid file type!</p>", 400


def convert_jpg_to_svg(img):
    max_size = (500, 500)  
    img.thumbnail(max_size)

    width, height = img.size
    svg_content = f'<svg width="{width}" height="{height}" xmlns="http://www.w3.org/2000/svg">'
    
    for y in range(height):
        for x in range(width):
            r, g, b, a = img.getpixel((x, y))
            if a > 0:  # Ignore fully transparent pixels
                color = f"#{r:02x}{g:02x}{b:02x}"
                svg_content += f'<rect x="{x}" y="{y}" width="1" height="1" fill="{color}" />'
    
    svg_content += "</svg>"
    return svg_content


@app.route("/converted-files", methods=["GET"])
def converted_files():
    files = list(mongo_db.files.find({}, {"filename": 1, "file_id": 1}))
    return render_template("converted.html", files=files)


def cleanup_gridfs():
    expired_files = mongo_db.files.find({"uploaded_at": {"$lt": datetime.utcnow() - timedelta(days=7)}})
    
    for file in expired_files:
        fs.delete(ObjectId(file["file_id"])) 
        mongo_db.files.delete_one({"_id": file["_id"]}) 


scheduler = BackgroundScheduler()
scheduler.add_job(cleanup_gridfs, "interval", hours=1)
scheduler.start()


mongo_db.files.create_index("uploaded_at", expireAfterSeconds=7*24*60*60)  

if __name__ == "__main__":
    app.run(debug=True)
